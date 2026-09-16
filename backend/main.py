import asyncio, io, json, math, os, secrets, time, zipfile
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Literal
from fastapi import FastAPI, HTTPException, Depends, Request, Response, UploadFile, File, Form
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from .knowledge_graph import kg
from pydantic import BaseModel, Field
from PIL import Image, ImageDraw
from . import store as s
from .schemas import MQTTConfig, Command, Question, Record, Device
from .mqtt_service import service
from .mqtt_config import read as read_mqtt_config
from .vision import vision, ModelUnavailable
from . import sessions
from .remote_camera import CameraFrameRequest, fetch_frame, probe
from . import camera_config
from .local_ai import local_ai, AIUnavailable
from . import camera_receiver
DEMO=os.getenv('HYDRO_DEMO','1')=='1'
ROLES={'teacher':'教师','student':'学生','admin':'管理员','parent':'家长'}
s.init();sessions.init();s.seed()
async def ticker():
 n=0
 while True:
  await asyncio.sleep(2);n+=1
  for d in s.allof('devices'):
   if d.get('source')=='demo' and DEMO and not s.REAL_ONLY:
    vals={k:round(v+(0 if k=='ph' else math.sin(n/30)*v*.008),2) for k,_,_,v in s.METRICS}
    s.add_telemetry(d['id'],d['batch'],vals)
   if n%5==0:
    from .control import evaluate_device
    evaluate_device(d)
  for c in s.allof('commands'):
   if c['status']=='sent_unconfirmed' and time.time()-datetime.fromisoformat(c['created_at']).timestamp()>(read_mqtt_config()['protocol']['result_timeout_seconds'] if c.get('protocol')=='tomato_v1_1' else 30):
    s.patch('commands',c['id'],{'status':'timeout'})
    if c.get('protocol')!='tomato_v1_1':s.patch('devices',c['device'],{'mode':'manual'})
    s.log('控制回执超时','未收到设备结果，请核对设备状态；设备 AI 模式保持不变。' if c.get('protocol')=='tomato_v1_1' else '已暂停自动策略，请检查设备后重新启用。','warning',device=c['device'])
@asynccontextmanager
async def lifespan(app):
 if os.getenv('YOLO_MODEL_PATH'):
  try:await asyncio.to_thread(vision.load)
  except ModelUnavailable:s.log('模型加载失败','识别服务不可用，请检查后端权重与运行依赖。','error')
 try:kg.initialize()
 except RuntimeError as exc:s.log('本地知识图谱加载失败',str(exc),'error')
 for d in s.allof('devices'):
  if d.get('mode')=='auto' and d.get('protocol')!='tomato_v1_1':s.patch('devices',d['id'],{'mode':'manual'})
 task=asyncio.create_task(ticker());mqtt_task=asyncio.create_task(service.maintain())
 try:yield
 finally:
  task.cancel();mqtt_task.cancel()
  await asyncio.gather(task,mqtt_task,return_exceptions=True)
  await asyncio.to_thread(service.disconnect)
app=FastAPI(title='多源感知水培智控系统 API',version='1.0.0',lifespan=lifespan)
app.include_router(camera_receiver.router)
@app.exception_handler(AIUnavailable)
async def ai_error(request,exc):
 return JSONResponse(status_code=503,content={'detail':str(exc)})
def actor(request:Request):
 u=sessions.get(request.cookies.get('hydro_session'),DEMO)
 if not u:raise HTTPException(401,'登录已失效，请重新登录后继续；已拍照片可在登录后重试')
 return u
def editor(u=Depends(actor)):
 if u['role']=='parent':raise HTTPException(403,'家长角色仅可查看成果')
 return u
def teacher(u=Depends(actor)):
 if u['role'] not in ['teacher','admin']:raise HTTPException(403,'此操作需要教师或管理员权限')
 return u
def admin(u=Depends(actor)):
 if u['role']!='admin':raise HTTPException(403,'连接配置需要管理员权限')
 return u
def obj(kind,id):
 v=s.get(kind,id)
 if v is None:raise HTTPException(404,'记录不存在')
 return v
def scoped(kind,u):
 rows=s.allof(kind)
 if u['role'] in ['student','parent'] and kind in ['records','submissions','favorites','inquiries']:
  rows=[x for x in rows if x.get('owner')=='student' or x.get('owner')==u['role']]
 return rows
@app.get('/api/health')
def health():return {'ok':True,'demo':DEMO,'real_only':s.REAL_ONLY,'model_configured':vision.status()['configured'],'model':vision.status()}
class Login(BaseModel):
 role:Literal['teacher','student','admin','parent']='teacher';password:str=''
@app.post('/api/session')
def login(v:Login,response:Response):
 if not DEMO:
  expected=os.getenv('HYDRO_'+v.role.upper()+'_PASSWORD')
  if not expected or not secrets.compare_digest(v.password,expected):raise HTTPException(401,'账户或密码不正确')
 token,u=sessions.create(v.role,ROLES[v.role],DEMO)
 response.set_cookie('hydro_session',token,httponly=True,samesite='strict',max_age=28800)
 return {**u,'demo':DEMO}
@app.get('/api/session')
def session(u=Depends(actor)):return {**u,'demo':DEMO}
@app.delete('/api/session')
def logout(request:Request,response:Response):
 sessions.delete(request.cookies.get('hydro_session'));response.delete_cookie('hydro_session');return {'ok':True}
@app.get('/api/devices')
def devices(u=Depends(actor)):
 out=[]
 for d in s.allof('devices'):
  t=s.telemetry(d['id'],1);latest=t[-1] if t and t[-1]['source']==d['source'] and t[-1]['batch']==d['batch'] else None
  config=service.config()
  tomato=d['source']=='mqtt' and config['protocol']=='tomato_v1_1' and d['id']==config['device_id']
  stale=read_mqtt_config()['protocol']['telemetry_stale_seconds'] if tomato else 15
  online=service.device_online(d) if tomato else bool(latest and time.time()-datetime.fromisoformat(latest['ts']).timestamp()<stale)
  unavailable=[key for key,_,_,_ in s.METRICS if latest['values'].get(key) is None] if latest else []
  out.append({**d,'online':online,'latest':latest,'unavailable_metrics':unavailable,'metric_units':{'level':'cm','ec':'mS/cm'} if tomato else d.get('metric_units',{}),'supported_actuators':['pump','light'] if tomato else ['pump','light','fan','mist'],'hardware_mode':tomato})
 return sorted(out,key=lambda d:d['id'])
@app.post('/api/devices')
def add_device(v:Device,u=Depends(teacher)):
 if s.REAL_ONLY and v.source!='mqtt':raise HTTPException(422,'仅允许真实 MQTT 设备')
 if s.get('devices',v.id):raise HTTPException(409,'设备编号已存在')
 return s.put('devices',{**v.model_dump(),'mode':'manual','actuators':dict(pump=False,light=False,fan=False,mist=False),'thresholds':{'ph_min':5.8,'ph_max':6.5,'temp_max':32,'level_min':20}})
@app.put('/api/devices/{id}')
def edit_device(id:str,v:Device,u=Depends(teacher)):
 obj('devices',id)
 if s.REAL_ONLY and v.source!='mqtt':raise HTTPException(422,'仅允许真实 MQTT 设备')
 if id!=v.id:raise HTTPException(422,'设备编号不可变更')
 return s.patch('devices',id,v.model_dump())
class Mode(BaseModel):
 mode:Literal['manual','auto'];confirmed:bool=False
@app.post('/api/devices/{id}/mode')
def mode(id:str,v:Mode,u=Depends(teacher)):
 d=obj('devices',id)
 if d['source']=='mqtt' and service.config()['protocol']=='tomato_v1_1':
  if v.mode=='auto' and not v.confirmed:raise HTTPException(422,'需确认启用设备 AI 智控')
  try:return service.protocol_command(id,'08',{'value':1 if v.mode=='auto' else 0})
  except ValueError as e:raise HTTPException(409,str(e))
 if v.mode=='auto':
  if not v.confirmed:raise HTTPException(422,'启用自动策略前必须确认规则')
  if d['source']=='mqtt' and not service.connected:raise HTTPException(409,'请先连接 MQTT')
  latest=s.telemetry(id,1)
  if not latest or latest[-1]['source']!=d['source'] or time.time()-datetime.fromisoformat(latest[-1]['ts']).timestamp()>15:raise HTTPException(409,'需要新鲜遥测数据才能启用自动策略')
 s.log('运行模式切换',v.mode,device=id)
 return s.patch('devices',id,{'mode':v.mode})
@app.get('/api/telemetry/{id}')
def telemetry(id:str,period:Literal['hour','day','week']='day',u=Depends(actor)):
 obj('devices',id)
 duration={'hour':3600,'day':86400,'week':604800}[period];bucket=60 if period=='hour' else 3600
 points=s.aggregate(id,duration,bucket)
 rows=s.telemetry(id,1)
 return {'points':points,'latest':rows[-1] if rows else None,'aggregation':'minute' if bucket==60 else 'hour'}

@app.post('/api/devices/{id}/snapshot')
def snapshot(id:str,u=Depends(editor)):
 d=obj('devices',id);rows=s.telemetry(id,1)
 if not rows or rows[-1]['source']!=d['source']:raise HTTPException(409,'当前数据通道尚无遥测，无法保存快照')
 missing=[label for key,label,_,_ in s.METRICS if rows[-1]['values'].get(key) is None]
 content='保存设备上报值；暂无读数：'+'、'.join(missing) if missing else '七项参数完整留存'
 return s.put('records',{'title':'环境数据快照','content':content,'type':'数据快照','device':id,'batch':d['batch'],'owner':u['role'],'photos':[],'values':rows[-1]['values'],'source':rows[-1]['source'],'telemetry_at':rows[-1]['ts'],'units':rows[-1].get('units',{})})
@app.get('/api/mqtt')
def mqtt_config(u=Depends(actor)):
 return {'config':service.config(True),**service.status()}
@app.put('/api/mqtt')
def save_mqtt(v:MQTTConfig,u=Depends(actor)):
 d=v.model_dump()
 if d['telemetry_topic']==d['ack_topic']:raise HTTPException(422,'遥测主题和回执主题不能相同')
 if not d['password']:d['password']=service.config()['password']
 service.disconnect();service.save_config(d);s.log('MQTT 配置已保存','配置已更新，连接已断开，需要重新测试或连接。')
 return service.config(True)
@app.post('/api/mqtt/{action}')
def mqtt_action(action:Literal['connect','test','disconnect'],u=Depends(actor)):
 try:
  if action=='disconnect':service.disconnect()
  else:
   service.start()
   if action=='test':service.disconnect()
 except ValueError as e:raise HTTPException(400,str(e))
 return {'connected':service.connected,'message':'连接测试通过' if action=='test' else ('已连接' if service.connected else '已断开')}
class ProtocolCommand(BaseModel):
 cmd:str;data:dict;confirmed:bool=False
@app.post('/api/devices/{id}/protocol-commands')
def protocol_command(id:str,v:ProtocolCommand,u=Depends(editor)):
 obj('devices',id)
 if v.cmd=='08':teacher(u)
 if not v.confirmed:raise HTTPException(422,'下发协议命令前必须确认')
 try:return service.protocol_command(id,v.cmd,v.data)
 except ValueError as e:raise HTTPException(409,str(e))
@app.post('/api/devices/{id}/commands')
def send_command(id:str,v:Command,u=Depends(editor)):
 d=obj('devices',id)
 if not v.confirmed:raise HTTPException(422,'请确认设备和操作后再发送')
 if d['source']=='demo':
  if s.REAL_ONLY:raise HTTPException(409,'真实数据模式禁止模拟控制')
  cmd=s.put('commands',{'device':id,'actuator':v.actuator,'state':v.state,'status':'simulated'})
  a=d['actuators'];a[v.actuator]=v.state;s.patch('devices',id,{'actuators':a})
  s.log('模拟设备控制',f'{v.actuator} → {v.state}（未发送到真实设备）',device=id,command=cmd);return cmd
 try:return service.command(id,v.actuator,v.state)
 except ValueError as e:raise HTTPException(409,str(e))
@app.get('/api/commands')
def commands(u=Depends(actor)):return s.allof('commands')[:100]
@app.get('/api/alerts')
def alerts(u=Depends(actor)):
 out=[]
 for d in devices(u):
  if not d['latest']:continue
  vals=d['latest']['values'];th=d['thresholds']
  level_min=th.get('level_min_cm',0) if d.get('metric_units',{}).get('level')=='cm' else th['level_min']
  for key,metric,limit,below,title in [('ph_low','ph',th['ph_min'],True,'pH 偏低'),('ph_high','ph',th['ph_max'],False,'pH 偏高'),('level','level',level_min,True,'液位不足'),('temp','air_temp',th['temp_max'],False,'空气温度过高')]:
   value=vals.get(metric)
   if value is not None and (value<limit if below else value>limit):
    id=d['id']+'-'+key;r=s.get('alert_resolutions',id)
    out.append({'id':id,'device':d['id'],'title':title,'resolved':bool(r),'resolution':r,'values':vals,'source':d['source'],'advice':'核对传感器读数，记录现场照片并检查设备状态。'})
 return out
@app.post('/api/alerts/{id}/resolve')
def resolve(id:str,u=Depends(editor)):
 if not any(a['id']==id for a in alerts(u)):raise HTTPException(404,'预警不存在')
 return s.put('alert_resolutions',{'id':id,'owner':u['role'],'note':'已人工标记处理；阈值异常仍会显示。'})
@app.get('/api/knowledge/status')
def knowledge_status(u=Depends(actor)):return kg.status()
@app.get('/api/ai/status')
def ai_status(u=Depends(actor)):return local_ai.status()
@app.get('/api/knowledge/graph')
def knowledge_overview(q:str='',limit:int=50,u=Depends(actor)):
 items=kg.all(q);limit=max(1,min(limit,100))
 return {'graph':kg.graph(items[:limit]),'items':items[:limit],'total':len(items),'limit':limit}
@app.get('/api/knowledge')
def knowledge(q:str='',u=Depends(actor)):return kg.all(q)
def graph(items):return kg.graph(items)
def retrieve(q):return kg.retrieve(q)
@app.post('/api/ask')
def ask(v:Question,u=Depends(actor)):
 environment=None;batch=''
 if v.device:
  d=obj('devices',v.device);batch=d['batch'];samples=s.telemetry(v.device,1)
  if samples and samples[-1]['source']==d['source'] and samples[-1]['batch']==batch:
   sample=samples[-1];age=max(0,time.time()-datetime.fromisoformat(sample['ts']).timestamp())
   environment={k:sample[k] for k in ('device','batch','ts','source','values','units')}
   environment.update(age_seconds=round(age),is_stale=age>45)
 items=retrieve(v.question)
 result={'question':v.question,'device':v.device,'batch':batch,'environment':environment,
         'items':items,'graph':graph(items),'mode':'python_graph','generation':None,
         'message':'' if items else '未检索到相关真实来源。请补充资料或换用资料中的关键词；本次未生成无来源的答案。'}
 if v.mode=='local_ai' and items:
  result['generation']=local_ai.generate(v.question,items,environment,v.model);result['mode']='local_ai'
 result=s.put('answers',{**result,'owner':u['role']})
 evidence={'knowledge':items,'answer_id':result['id'],'generation':result['generation'],'batch':batch}
 if environment:evidence['environment']=environment
 s.put('logs',dict(title='本地 AI 知识问答' if result['generation'] else '知识检索与图谱关联',
                  detail=v.question,level='info',device=v.device,sources=[k['id'] for k in items],
                  evidence=evidence,flagged=False))
 return result
class Knowledge(BaseModel):
 model_config={'str_strip_whitespace':True}
 crop:str=Field(default='番茄',min_length=1,max_length=100);environments:list[str]=Field(default_factory=list,max_length=12)
 title:str=Field(min_length=1,max_length=200);problem:str=Field(min_length=1,max_length=300);cause:str=Field(min_length=1,max_length=500);measure:str=Field(min_length=1,max_length=1000);content:str=Field(min_length=1,max_length=20000);source:str=Field(default='种植经验',min_length=1,max_length=1000);tags:list[str]=Field(default_factory=list,max_length=30)
@app.post('/api/knowledge')
def create_knowledge(v:Knowledge,u=Depends(teacher)):
 return kg.upsert({**v.model_dump(),'id':'KB-'+s.uid()[:6].upper(),'created_at':s.now(),'verified':False,'owner':u['role']})
@app.put('/api/knowledge/{id}')
def edit_knowledge(id:str,v:Knowledge,u=Depends(teacher)):
 old=next((k for k in kg.all() if k['id']==id),None)
 if old is None:raise HTTPException(404,'知识资料不存在')
 revision={**old,**v.model_dump(),'updated_at':s.now(),'verified':False,'owner':u['role']}
 revision['revisions']=(old.get('revisions',[])+[{k:value for k,value in old.items() if k!='revisions'}])
 return kg.upsert(revision)
@app.get('/api/summary/{id}')
def summary(id:str,u=Depends(actor)):
 d=obj('devices',id);today=datetime.now().astimezone().replace(hour=0,minute=0,second=0,microsecond=0).timestamp()
 events=[c for c in s.allof('commands') if c['device']==id and c['status'] in ('simulated','acknowledged')]
 water=sum(1 for c in events if c['actuator']=='pump' and c['state'] and datetime.fromisoformat(c['created_at']).timestamp()>=today)
 lights=sorted([c for c in events if c['actuator']=='light'],key=lambda c:c['created_at']);on=False;last=today;duration=0
 for c in lights:
  at=datetime.fromisoformat(c.get('ack_at',c['created_at'])).timestamp()
  if at<today:on=c['state'];continue
  if on:duration+=at-last
  last=at;on=c['state']
 if on:duration+=time.time()-last
 count=sum(1 for r in s.allof('recognitions') if r['device']==id and datetime.fromisoformat(r['created_at']).timestamp()>=today)
 return {'water_count':water,'light_hours':round(duration/3600,2),'recognition_count':count,'scope':'已记录的模拟执行或设备回执，不估算未记录动作'}
@app.get('/api/achievements/me')
def achievements(u=Depends(actor)):
 subs=[x for x in scoped('submissions',u) if x.get('owner')==('student' if u['role']=='parent' else u['role'])]
 records=scoped('records',u);inquiries=scoped('inquiries',u)
 points=sum(20 if x['status']=='evaluated' else 10 if x['status']=='submitted' else 0 for x in subs)
 return {'points':points,'badges':[{'name':'观察新芽','icon':'seedling-line','earned':bool(records)},{'name':'知识探索者','icon':'book-open-line','earned':bool(inquiries)},{'name':'实训践行者','icon':'medal-line','earned':any(x['status']=='evaluated' for x in subs)}]}
@app.get('/api/{kind}')
def list_objects(kind:Literal['records','inquiries','favorites','tasks','submissions','logs','recognitions'],device:str='',batch:str='',q:str='',u=Depends(actor)):
 return [x for x in scoped(kind,u) if (not device or x.get('device')==device) and (not batch or x.get('batch')==batch) and (not q or q in json.dumps(x,ensure_ascii=False))]
@app.post('/api/records')
def record(v:Record,u=Depends(editor)):
 obj('devices',v.device)
 for p in v.photos:
  if not s.get('photos',p):raise HTTPException(422,'照片不存在')
 return s.put('records',{**v.model_dump(),'owner':u['role']})
@app.post('/api/inquiries')
def inquiry(v:Record,u=Depends(editor)):
 if v.answer_id:
  answer=owned_answer(v.answer_id,u)
  return s.put('inquiries',{'title':answer['question'],'content':answer_text(answer),'type':'知识探究',
    'device':answer['device'],'batch':answer['batch'],'sources':[k['id'] for k in answer['items']],
    'answer':answer,'owner':u['role']})
 return s.put('inquiries',{**v.model_dump(),'owner':u['role']})
def owned_answer(id,u):
 answer=obj('answers',id)
 if answer['owner']!=u['role']:raise HTTPException(403,'不能保存其他账户的问答')
 return answer
def answer_text(answer):
 generation=answer.get('generation')
 if generation:
  text='\n\n'.join(x['text']+' ['+', '.join(x['source_ids'])+']' for x in generation['statements'])
  if generation.get('follow_up'):text+='\n\n待补充：'+'；'.join(generation['follow_up'])
  return text
 return '\n\n'.join(k['content']+' ['+k['id']+']' for k in answer['items']) or answer['message']
@app.post('/api/favorites')
def favorite(v:Question,u=Depends(editor)):
 if v.answer_id:
  answer=owned_answer(v.answer_id,u)
  return s.put('favorites',{'title':answer['question'],'content':answer['items'],
      'sources':[k['id'] for k in answer['items']],'graph':answer['graph'],
      'answer':answer,'owner':u['role']})
 items=retrieve(v.question)
 if not items:raise HTTPException(422,'没有可收藏的答案')
 return s.put('favorites',{'title':v.question,'content':items,'sources':[k['id'] for k in items],'graph':graph(items),'owner':u['role']})
class Flag(BaseModel):flagged:bool=True
@app.patch('/api/logs/{id}')
def flag(id:str,v:Flag,u=Depends(teacher)):
 obj('logs',id);return s.patch('logs',id,v.model_dump())
class Task(BaseModel):
 title:str=Field(min_length=1,max_length=200);description:str=Field(min_length=1,max_length=5000);deadline:str;hours:int=Field(default=2,ge=1,le=100)
@app.post('/api/tasks')
def create_task(v:Task,u=Depends(teacher)):return s.put('tasks',{**v.model_dump(),'status':'published','teacher':u['name']})
@app.post('/api/tasks/{id}/claim')
def claim(id:str,u=Depends(editor)):
 obj('tasks',id)
 old=next((x for x in s.allof('submissions') if x['task_id']==id and x['owner']==u['role']),None)
 if old:return old
 return s.put('submissions',{'task_id':id,'owner':u['role'],'status':'claimed','content':'','photos':[],'records':[],'evaluation':''})
class Submission(BaseModel):content:str=Field(min_length=1,max_length=20000);photos:list[str]=[];records:list[str]=[]
@app.post('/api/tasks/{id}/submit')
def submit(id:str,v:Submission,u=Depends(editor)):
 sub=claim(id,u)
 for p in v.photos:
  if not s.get('photos',p):raise HTTPException(422,'照片不存在')
 for r in v.records:
  rec=obj('records',r)
  if rec.get('owner')!=u['role'] and u['role'] not in ['teacher','admin']:raise HTTPException(403,'不能提交他人的记录')
 return s.patch('submissions',sub['id'],{**v.model_dump(),'status':'submitted','submitted_at':s.now()})
class Evaluation(BaseModel):evaluation:str=Field(min_length=1,max_length=5000);score:int=Field(ge=0,le=100)
@app.post('/api/submissions/{id}/evaluate')
def evaluate(id:str,v:Evaluation,u=Depends(teacher)):
 sub=obj('submissions',id)
 if sub['status'] not in ['submitted','evaluated']:raise HTTPException(409,'成果尚未提交')
 return s.patch('submissions',id,{**v.model_dump(),'status':'evaluated','reviewer':u['name']})
@app.post('/api/camera/frame')
async def camera_frame(v:CameraFrameRequest,response:Response,u=Depends(editor)):
 response.headers['Cache-Control']='no-store'
 return await fetch_frame(v.url)
@app.post('/api/camera/probe')
async def camera_probe(v:CameraFrameRequest,response:Response,u=Depends(editor)):
 response.headers['Cache-Control']='no-store'
 return await probe(v.url)
class CameraAddress(BaseModel):
 url:str=Field(min_length=1,max_length=512)
@app.get('/api/camera/config')
def camera_address(u=Depends(actor)):
 return camera_config.read()
@app.get('/api/camera/receiver')
def camera_receiver_status(response:Response,u=Depends(actor)):
 response.headers['Cache-Control']='no-store'
 return camera_receiver.metadata(camera_receiver.latest())
@app.put('/api/camera/config')
def save_camera_address(v:CameraAddress,u=Depends(editor)):
 saved=camera_config.save(v.url)
 s.log('摄像头地址已更新',f"远程摄像头默认地址保存为 {saved['url']}，页面与后端都会请求该地址。")
 return saved

@app.post('/api/photos')
async def photo(file:UploadFile=File(...),device:str=Form('HY-001'),batch:str=Form('2026-A'),capture_source:Literal['upload','local_camera','remote_camera']=Form('upload'),u=Depends(editor)):
 obj('devices',device);data=await file.read(12*1024*1024+1)
 if len(data)>12*1024*1024:raise HTTPException(413,'图片不能超过 12 MB')
 try:
  img=Image.open(io.BytesIO(data));img.verify();img=Image.open(io.BytesIO(data)).convert('RGB');img.thumbnail((2400,2400))
 except Exception:raise HTTPException(422,'请上传有效的图片')
 id=s.uid();path=s.ROOT/'photos'/f'{id}.jpg';img.save(path,quality=92)
 return s.put('photos',{'id':id,'device':device,'batch':batch,'owner':u['role'],'url':f'/api/photos/{id}','name':Path(file.filename or 'photo').name,'capture_source':capture_source})
@app.get('/api/photos/{id}')
def get_photo(id:str,u=Depends(actor)):
 p=obj('photos',id);return FileResponse(s.ROOT/'photos'/f'{id}.jpg',media_type='image/jpeg')
def recognition_summary(detections):
 distribution={'未成熟':0,'半成熟':0,'成熟':0}
 for detection in detections:
  label=detection['label'];distribution[label]=distribution.get(label,0)+1
 return {'fruit_count':len(detections),'maturity_distribution':distribution,'average_confidence':round(sum(d['confidence'] for d in detections)/len(detections),4) if detections else None}

class Recognition(BaseModel):
 photo_id:str
 demo:bool=False
 confidence_threshold:float=Field(default=0.65,ge=0.1,le=0.95)
@app.post('/api/recognitions')
def recognize(v:Recognition,u=Depends(editor)):
 p=obj('photos',v.photo_id);path=s.ROOT/'photos'/f'{v.photo_id}.jpg';detections=[];model=os.getenv('YOLO_MODEL_PATH');model_info=None
 with Image.open(path) as original:img=original.convert('RGB')
 if v.demo:
  if s.REAL_ONLY:raise HTTPException(403,'真实数据模式禁止演示识别')
  if not DEMO:raise HTTPException(403,'当前环境未开启演示')
  w,h=img.size
  detections=[{'label':'成熟','confidence':.94,'box':[int(w*.22),int(h*.2),int(w*.44),int(h*.55)]},{'label':'半成熟','confidence':.86,'box':[int(w*.56),int(h*.35),int(w*.76),int(h*.7)]}]
 elif model:
  try:
   detections,model_info=vision.predict(str(path),confidence_threshold=v.confidence_threshold)
  except ModelUnavailable as e:s.log('模型推理失败',str(e),'error');raise HTTPException(503,str(e))
 else:raise HTTPException(503 if s.REAL_ONLY else 409,'尚未配置 YOLO 模型，请联系管理员配置后端模型权重')
 detections=[d for d in detections if d['confidence']>=v.confidence_threshold]
 draw=ImageDraw.Draw(img)
 for i,d in enumerate(detections):draw.rectangle(d['box'],outline='#24ab60',width=4);draw.text((d['box'][0]+3,d['box'][1]+3),f'{i+1} {d["confidence"]:.0%}',fill='#24ab60')
 aid=s.uid();img.save(s.ROOT/'photos'/f'{aid}.jpg');s.put('photos',{**p,'id':aid,'name':'annotated.jpg','url':f'/api/photos/{aid}'})
 r=s.put('recognitions',{'device':p['device'],'batch':p['batch'],'original':p['url'],'annotated':f'/api/photos/{aid}','photo_ids':[v.photo_id,aid],'detections':detections,'summary':recognition_summary(detections),'mode':'demo' if v.demo else 'yolo','model':model_info,'owner':u['role'],'title':'成熟度识别记录'})
 s.put('records',{'title':'成熟度识别留存','content':json.dumps(detections,ensure_ascii=False),'type':'AI 识别','device':p['device'],'batch':p['batch'],'photos':[v.photo_id,aid],'owner':u['role'],'recognition':r['id'],'summary':r['summary']})
 s.log('成熟度识别完成','演示结果（非真实推理）' if v.demo else 'YOLO 推理完成',device=p['device'],recognition=r);return r

def export_rows(kind,device,u):
 if kind=='telemetry':return s.telemetry(device,20000)
 if kind=='devices':return devices(u)
 rows=[x for x in scoped(kind,u) if not device or x.get('device')==device or 'device' not in x]
 if kind=='recognitions':rows=[{**x,'summary':recognition_summary(x['detections'])} for x in rows]
 return rows
@app.get('/api/export/{kind}')
def export(kind:Literal['records','logs','telemetry','devices','recognitions','inquiries','submissions','favorites'],format:Literal['csv','xlsx','pdf','json','zip']='csv',device:str='',u=Depends(actor)):
 rows=export_rows(kind,device or ('HY-001' if kind=='telemetry' else ''),u)
 if format=='json':return Response(json.dumps(rows,ensure_ascii=False,indent=2),media_type='application/json',headers={'Content-Disposition':f'attachment; filename="{kind}.json"'})
 flat=[{k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in r.items()} for r in rows]
 cols=list(dict.fromkeys(k for r in flat for k in r)) or ['title']
 buf=io.BytesIO()
 if format=='csv':
  import csv
  out=io.StringIO();w=csv.DictWriter(out,fieldnames=cols);w.writeheader()
  for row in flat:
   # Prevent exported user text being executed as spreadsheet formulas.
   w.writerow({k:("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for k,v in row.items()})
  buf.write(out.getvalue().encode('utf-8-sig'));mime='text/csv'
 elif format=='xlsx':
  from openpyxl import Workbook
  wb=Workbook();ws=wb.active;ws.title='水培数据';ws.append(cols)
  for r in flat:
   ws.append([r.get(k,'') for k in cols])
   for cell in ws[ws.max_row]:
    if isinstance(cell.value,str):cell.data_type='s'
  wb.save(buf);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 elif format=='zip':
  with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
   z.writestr('records.json',json.dumps(rows,ensure_ascii=False,indent=2))
   photos=set(p for r in rows for p in r.get('photos',r.get('photo_ids',[])))
   for pid in photos:
    path=s.ROOT/'photos'/f'{pid}.jpg'
    if path.exists():z.write(path,'photos/'+path.name)
  mime='application/zip'
 else:
  from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage
  from reportlab.lib.styles import getSampleStyleSheet
  from reportlab.pdfbase import pdfmetrics
  from reportlab.pdfbase.cidfonts import UnicodeCIDFont
  from xml.sax.saxutils import escape
  pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'));styles=getSampleStyleSheet()
  for st in styles.byName.values():st.fontName='STSong-Light';st.wordWrap='CJK'
  styles['BodyText'].fontSize=10;styles['BodyText'].leading=16
  flow=[Paragraph('多源感知水培智控系统 '+{'records':'生长档案报告','logs':'AI 复盘报告','recognitions':'采摘报告','devices':'设备对比报告','inquiries':'探究报告'}.get(kind,'数据报告'),styles['Title']),Paragraph('导出时间：'+datetime.now().strftime('%Y-%m-%d %H:%M')+' · 共 '+str(len(rows))+' 条记录',styles['BodyText']),Spacer(1,16)]
  if kind in ('records','telemetry'):
   from reportlab.graphics.shapes import Drawing, String
   from reportlab.graphics.charts.lineplots import LinePlot
   from reportlab.graphics.widgets.markers import makeMarker
   from reportlab.lib.colors import HexColor
   target=device or (rows[0].get('device') if rows else None)
   series=s.aggregate(target,86400,3600) if target else []
   if series:
    flow.append(Paragraph('最近 24 小时环境变化（小时聚合）',styles['Heading2']))
    for key,label,unit,_ in s.METRICS:
     chart_device=s.get('devices',target) or {}
     unit=(chart_device.get('metric_units',{}) if chart_device.get('source')=='mqtt' else {}).get(key,unit)
     # Split at unavailable readings instead of inventing zeroes or bridging gaps.
     segments=[];segment=[]
     for i,point in enumerate(series):
      value=point['values'].get(key)
      if value is None:
       if segment:segments.append(segment);segment=[]
      else:segment.append((i,value))
     if segment:segments.append(segment)
     if not segments:
      flow.append(Paragraph(label+'：暂无读数',styles['BodyText']));continue
     drawing=Drawing(440,135);chart=LinePlot();chart.x=40;chart.y=22;chart.width=375;chart.height=90
     chart.data=segments
     for i in range(len(segments)):
      chart.lines[i].strokeColor=HexColor('#347854');chart.lines[i].strokeWidth=1.5;chart.lines[i].symbol=makeMarker('FilledCircle');chart.lines[i].symbol.size=3
     chart.xValueAxis.valueMin=0;chart.xValueAxis.valueMax=max(1,len(series)-1);chart.xValueAxis.labelTextFormat=lambda v: str(int(v))
     chart.yValueAxis.labels.fontSize=8;chart.xValueAxis.labels.fontSize=8;drawing.add(chart)
     drawing.add(String(40,120,label+' ('+unit+')',fontName='STSong-Light',fontSize=10));flow.append(drawing)
    flow.append(Paragraph('横轴为连续小时时间点，起始时间 '+series[0]['ts'][:16],styles['BodyText']))
  labels={'created_at':'记录时间','device':'设备编号','batch':'茬次','type':'记录类型','content':'内容','values':'环境参数','source':'数据来源','mode':'识别模式','evidence':'关联证据','summary':'识别统计','detections':'识别结果','owner':'记录人','detail':'详情','sources':'知识来源','evaluation':'教师评价','score':'分数','title':'标题','ts':'采集时间','name':'设备名称','group':'班级分组','status':'状态','actuators':'执行器状态'}
  for r in rows[:500]:
   flow.append(Paragraph(escape(str(r.get('title',r.get('ts',r.get('id',''))))),styles['Heading2']))
   for k,v in r.items():
    if k in ('photos','photo_ids','original','annotated','id'):continue
    if k=='summary':
     v='果实数量：'+str(v['fruit_count'])+'；成熟度分布：'+'，'.join(f'{label} {count}' for label,count in v['maturity_distribution'].items())+'；平均置信度：'+(f"{v['average_confidence']:.1%}" if v['average_confidence'] is not None else '无检测目标')
    flow.append(Paragraph(escape(labels.get(k,k)+'：'+(json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v))),styles['BodyText']))
   for pid in r.get('photos',r.get('photo_ids',[])):
    p=s.ROOT/'photos'/f'{pid}.jpg'
    if p.exists():
     im=Image.open(p);w,h=im.size;flow.append(PDFImage(str(p),width=min(260,420*w/h),height=min(420,260*h/w)))
   flow.append(Spacer(1,12))
  SimpleDocTemplate(buf).build(flow);mime='application/pdf'
 buf.seek(0);return StreamingResponse(buf,media_type=mime,headers={'Content-Disposition':f'attachment; filename="{kind}.{format}"'})
